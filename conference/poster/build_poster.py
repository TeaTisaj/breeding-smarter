"""Draft poster (90 x 120 cm portrait PPTX) from the content pack (P11). Layout sketch until
the official BioMedVetMech template arrives: then copy the blocks into the template.

Usage: uv run python conference/build_texts.py && uv run python conference/build_figures.py
       && uv run python conference/poster/build_poster.py
"""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Cm, Pt

HERE = Path(__file__).resolve().parent
BLUE = RGBColor(0x00, 0x72, 0xB2)


def blocks(md: str) -> tuple[str, str, dict[str, str], str]:
    lines = [ln for ln in md.splitlines() if not ln.startswith("<!--")]
    title = lines[0].lstrip("# ").strip()
    authors = lines[1].strip()
    sections: dict[str, str] = {}
    current = ""
    refs = ""
    for ln in lines[2:]:
        if ln.startswith("## "):
            current = ln[3:].strip()
            sections[current] = ""
        elif ln.startswith("Refs:"):
            refs = ln
        elif current and ln.strip():
            sections[current] += re.sub(r"^- ", "• ", ln.strip()) + "\n"
    return title, authors, sections, refs


def textbox(slide, x, y, w, h, text, size, bold=False, colour=None):  # type: ignore[no-untyped-def]
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, para in enumerate(text.strip().split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = para
        p.font.size = Pt(size)
        p.font.bold = bold
        if colour is not None:
            p.font.color.rgb = colour
    return tb


def main() -> Path:
    title, authors, sec, refs = blocks((HERE / "text.md").read_text())
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(90), Cm(120)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    band = s.shapes.add_shape(1, Cm(0), Cm(0), Cm(90), Cm(14))
    band.fill.solid()
    band.fill.fore_color.rgb = BLUE
    band.line.fill.background()
    textbox(s, 3, 1.5, 84, 8, title, 66, True, RGBColor(0xFF, 0xFF, 0xFF))
    textbox(s, 3, 10, 84, 3, authors + "  [FER logo]", 32, False, RGBColor(0xFF, 0xFF, 0xFF))
    y = 16.0
    for name, h in [("The problem", 8), ("Why now", 4.5), ("What we built", 7.5)]:
        textbox(s, 3, y, 84, 2, name, 44, True, BLUE)
        textbox(s, 3, y + 2.3, 84, h - 2.3, sec.get(name, ""), 30)
        y += h
    textbox(s, 3, y, 84, 2.5, "[pipeline diagram: insert conference/poster/pipeline.svg]", 24)
    y += 3
    textbox(s, 3, y, 84, 2, "Recorded COI hides inbreeding", 44, True, BLUE)
    textbox(s, 3, y + 2.3, 84, 2.5, sec.get("Recorded COI hides inbreeding", ""), 30)
    pic = s.shapes.add_picture(
        str(HERE / "fig4_recorded_vs_true_f_poster.png"), Cm(3), Cm(y + 5), width=Cm(84)
    )
    y += 5 + pic.height.cm + 1
    textbox(s, 3, y, 41, 2, "Does it work?", 44, True, BLUE)
    textbox(s, 3, y + 2.3, 41, 9, sec.get("Does it work?", ""), 28)
    headline = HERE / "figG_coordination_poster.png"  # P12; falls back to the P8 figure
    if not headline.exists():
        headline = HERE / "figA_realised_inbreeding_poster.png"
    a = s.shapes.add_picture(str(headline), Cm(3), Cm(y + 11.5), width=Cm(41))
    textbox(s, 46, y, 41, 2, "Explainable matches", 44, True, BLUE)
    b = s.shapes.add_picture(str(HERE / "scorecard.png"), Cm(46), Cm(y + 2.5), width=Cm(41))
    y += max(11.5 + a.height.cm, 2.5 + b.height.cm) + 1
    textbox(s, 3, y, 41, 2, "Honesty box", 40, True, BLUE)
    textbox(s, 3, y + 2.3, 41, 6, sec.get("Honesty box", ""), 28)
    textbox(s, 46, y, 41, 2, "Partner with us", 40, True, BLUE)
    textbox(s, 46, y + 2.3, 41, 6, sec.get("Partner with us", "") + "[QR: demo, DOI, contact]", 28)
    y += 9
    textbox(s, 3, 116, 84, 3, refs, 20)
    out = HERE / "poster_draft.pptx"
    prs.save(out)
    print(f"wrote {out} (layout sketch; y used {y:.0f} cm of 120)")
    return out


if __name__ == "__main__":
    main()
